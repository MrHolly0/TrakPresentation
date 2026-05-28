from __future__ import annotations

import json
import os
from pathlib import Path

import docx
import openpyxl
from pypdf import PdfReader


ROOT = Path("/Users/holly/Desktop/На проверку")
OUT = ROOT / "07_Отчёт" / "materials_extracted.json"


def text_from_docx(path: Path) -> dict:
    doc = docx.Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append(rows)
    return {"paragraphs": paragraphs, "tables": tables}


def text_from_xlsx(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    sheets = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            vals = []
            for cell in row:
                val = cell.value
                if val is None:
                    vals.append("")
                else:
                    vals.append(str(val))
            while vals and vals[-1] == "":
                vals.pop()
            if any(v != "" for v in vals):
                rows.append(vals)
        sheets[ws.title] = rows
    return {"sheets": sheets}


def text_from_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append({"page": i, "text": text.strip()})
    return {"pages": pages}


def main() -> None:
    records = []
    for path in sorted(ROOT.glob("*/*")):
        if path.name.startswith(".") or path.is_dir():
            continue
        suffix = path.suffix.lower()
        rel = str(path.relative_to(ROOT))
        try:
            if suffix == ".docx":
                content = text_from_docx(path)
            elif suffix == ".xlsx":
                content = text_from_xlsx(path)
            elif suffix == ".pdf":
                content = text_from_pdf(path)
            else:
                content = {"note": "not extracted"}
            records.append({"file": rel, "type": suffix, "content": content})
        except Exception as exc:
            records.append({"file": rel, "type": suffix, "error": repr(exc)})
    OUT.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
