from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path("/Users/holly/Desktop/На проверку")


def replace_in_docx(path: Path, replacements: dict[str, str]) -> int:
    doc = Document(path)
    count = 0
    containers = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                containers.extend(cell.paragraphs)
    for p in containers:
        text = p.text
        new_text = text
        for old, new in replacements.items():
            if old in new_text:
                new_text = new_text.replace(old, new)
        if new_text != text:
            p.text = new_text
            count += 1
    if count:
        doc.save(path)
    return count


def rewrite_hr_task4(path: Path) -> None:
    doc = Document()
    doc.add_heading("Чек-лист технического скрининга senior ИБ-архитектора", level=1)
    doc.add_paragraph(
        "Чек-лист используется перед передачей кандидата клиенту. Цель — проверить не только наличие инструментов в резюме, "
        "но и глубину архитектурного опыта, способность объяснять риски бизнесу и готовность к переходу."
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Блок"
    hdr[1].text = "Что проверяем"
    hdr[2].text = "Вопрос / критерий"
    rows = [
        ("Опыт", "Senior-уровень в ИБ", "Какие зоны безопасности кандидат проектировал сам, а не только сопровождал?"),
        ("Архитектура", "Защитный контур", "Как выстроить сегментацию, IAM, мониторинг и реагирование для enterprise-системы?"),
        ("Нормативка", "187-ФЗ, регуляторика, внутренние политики", "С какими аудитами, проверками или требованиями регуляторов работал кандидат?"),
        ("Инциденты", "Практический опыт расследований", "Опишите сложный инцидент: причина, действия, выводы, изменения после."),
        ("Инструменты", "SIEM, DLP, IAM/PAM, EDR, сканеры", "Какие системы внедрял, администрировал или выбирал как архитектор?"),
        ("DevSecOps", "Безопасность в CI/CD", "Как встроить security checks в пайплайн без блокировки разработки?"),
        ("Коммуникация", "Работа с CTO, бизнесом и разработкой", "Как кандидат объясняет технический риск не техническому заказчику?"),
        ("Резюме", "Глубина роли", "Есть ли масштабы инфраструктуры, результаты и ответственность, а не только список инструментов?"),
        ("Мотивация", "Готовность к переходу", "Что должно измениться, чтобы кандидат реально принял оффер?"),
        ("Итог", "Решение по передаче клиенту", "Передавать / резерв / отказ; ключевые риски и условия оффера."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    doc.add_paragraph("Итоговый вердикт рекрутера: передавать клиенту / оставить в резерве / отказать.")
    doc.add_paragraph("Комментарий: указывать сильные стороны, риски, зарплатные ожидания и условия, при которых кандидат готов продолжить процесс.")
    doc.save(path)


def main() -> None:
    fixes = {}
    fixes["01_Аналитика/IT-рекрутинг 5 задание.docx"] = replace_in_docx(
        ROOT / "01_Аналитика/IT-рекрутинг 5 задание.docx",
        {"14,2": "16,1", "14.2": "16.1"},
    )
    fixes["01_Аналитика/IT-рекрутинг 1 задание.docx"] = replace_in_docx(
        ROOT / "01_Аналитика/IT-рекрутинг 1 задание.docx",
        {"14,2": "16,1", "14.2": "16.1", "?utm_source=chatgpt.com": ""},
    )
    fixes["01_Аналитика/IT-рекрутинг 2 задание.docx"] = replace_in_docx(
        ROOT / "01_Аналитика/IT-рекрутинг 2 задание.docx",
        {"?utm_source=chatgpt.com": ""},
    )
    fixes["01_Аналитика/IT-рекрутинг 6 задание.docx"] = replace_in_docx(
        ROOT / "01_Аналитика/IT-рекрутинг 6 задание.docx",
        {"?utm_source=chatgpt.com": ""},
    )
    fixes["05_Маркетинг/Маркетолог ПД.docx"] = replace_in_docx(
        ROOT / "05_Маркетинг/Маркетолог ПД.docx",
        {
            "МАРКЕКТИНГ": "МАРКЕТИНГ",
            "Нашли кандидата через личные рекомендации за 2 недели.": "Нашли и вывели кандидата на работу ровно за 21 день через личные рекомендации.",
            "за 2 недели": "за 21 день",
        },
    )
    rewrite_hr_task4(ROOT / "04_HR-процессы/HR-Таск 4.docx")
    fixes["04_HR-процессы/HR-Таск 4.docx"] = "rewritten"
    for file_name, result in fixes.items():
        print(file_name, result)


if __name__ == "__main__":
    main()
